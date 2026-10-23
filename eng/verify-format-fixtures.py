"""Independent original format-fixture checks; not production qualification.

Requires caller-provisioned CairoSVG, Pillow and PyMuPDF. No input resources are
fetched: only the original controlled generated SVG/PDF fixtures are rendered.
"""
import argparse
import hashlib
import json
from pathlib import Path
import io
import fitz
import cairosvg
from PIL import Image, ImageChops, ImageStat
import PIL


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def pdf_pixels(path):
    with fitz.open(path) as document:
        require(document.page_count == 1, f"{path.name}: expected one original vector page")
        pixel = document[0].get_pixmap(matrix=fitz.Matrix(96 / 72, 96 / 72), alpha=False)
        return Image.frombytes("RGB", (pixel.width, pixel.height), pixel.samples)


def spreadsheet_data(root):
    import openpyxl
    import zipfile
    import xml.etree.ElementTree as ET
    from odf.opendocument import load
    from odf.table import Table, TableCell
    from odf import teletype
    import importlib.metadata as metadata
    workbook = openpyxl.load_workbook(root / "table.xlsx", data_only=False)
    require(workbook.sheetnames == ["Data"], "XLSX sheet identity changed")
    sheet = workbook["Data"]
    require(sheet["A1"].value == "=SUM(A1:A2)" and sheet["A1"].data_type == "s", "Literal formula text became executable")
    require(str(sheet["A2"].value) == "12.5" and sheet["B2"].value is True, "XLSX scalar types changed")
    require(sheet["C2"].value.strftime("%Y-%m-%d") == "2026-10-20", "XLSX ISO date changed")
    require("A1:B1" in [str(r) for r in sheet.merged_cells.ranges], "XLSX merged rectangle changed")
    with zipfile.ZipFile(root / "table.xlsx") as package:
        for name in package.namelist():
            if name.startswith("xl/worksheets/"):
                xml = ET.fromstring(package.read(name))
                require(not xml.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f"), "Executable XLSX formula emitted")
    ods = load(str(root / "table.ods"))
    tables = ods.spreadsheet.getElementsByType(Table)
    require(len(tables) == 1 and tables[0].getAttribute("name") == "Data", "ODS table identity changed")
    cells = tables[0].getElementsByType(TableCell)
    require(cells[0].getAttribute("valuetype") == "string" and teletype.extractText(cells[0]) == "=SUM(A1:A2)", "ODS literal text changed")
    require(cells[0].getAttribute("numbercolumnsspanned") == "2", "ODS merge changed")
    require(all(cell.getAttribute("formula") is None for cell in cells), "Executable ODS formula emitted")
    xml = ET.parse(root / "table.xml")
    ns = {"s": "urn:schemas-microsoft-com:office:spreadsheet"}
    first = xml.find(".//s:Data", ns)
    require(first.text == "=SUM(A1:A2)" and first.get("{" + ns["s"] + "}Type") == "String", "SpreadsheetML literal type changed")
    require(not any("{" + ns["s"] + "}Formula" in cell.attrib for cell in xml.findall(".//s:Cell", ns)), "Executable SpreadsheetML formula emitted")
    return {"validators": {"openpyxl": openpyxl.__version__, "odfpy": metadata.version("odfpy"), "xml": "Python ElementTree"},
            "typedValuesMergesAndLiteralFormulas": "passed", "productionQualified": False,
            "artifacts": [{"file": name, "sha256": digest(root / name)} for name in ["table.xlsx", "table.ods", "table.xml"]]}


def presentation_data(root):
    from pptx import Presentation
    import pptx
    deck = Presentation(root / "raster-slides.pptx")
    require((deck.slide_width, deck.slide_height) == (914400, 914400), "PPTX common slide dimensions changed")
    require(len(deck.slides) == 3, "PPTX selected page count changed")
    for index, slide in enumerate(deck.slides):
        require(len(slide.shapes) == 1, "Raster profile emitted unexpected editable objects")
        picture = slide.shapes[0]
        image = Image.open(io.BytesIO(picture.image.blob))
        expected = 0 if index == 1 else 40
        require(image.getpixel((0, 0)) == (expected, 0, 0, 255), "PPTX page ordering or lossless RGB samples changed")
        require(picture.width == 914400 and picture.height == (457200 if index == 1 else 914400), "PPTX contain aspect changed")
        require(picture.top == (228600 if index == 1 else 0), "PPTX contain position changed")
    with fitz.open(root / "raster-slides.pdf") as pdf:
        require(pdf.page_count == 3 and all(page.rect.width == 72 and page.rect.height == 72 for page in pdf), "PPTX import changed page order/geometry")
    return {"validator": "python-pptx", "version": pptx.__version__, "orderedRasterGeometryAndPixels": "passed", "productionQualified": False,
            "artifacts": [{"file": name, "sha256": digest(root / name)} for name in ["raster-slides.pptx", "raster-slides.pdf"]]}


def html_data(root):
    from lxml import html, etree
    import base64
    fixed = html.fromstring((root / "fixed-pages.html").read_bytes())
    require(not fixed.xpath("//script|//iframe|//object|//a|//form"), "Active HTML content emitted")
    csp = fixed.xpath("//meta[@http-equiv='Content-Security-Policy']/@content")
    require(len(csp) == 1 and "default-src 'none'" in csp[0] and "img-src data:" in csp[0] and "base-uri 'none'" in csp[0], "Offline HTML policy changed")
    sections = fixed.xpath("//section")
    require([section.get("data-source-page") for section in sections] == ["1", "0", "1"], "HTML selected page order changed")
    images = []
    for index, section in enumerate(sections):
        values = section.xpath("./img/@src")
        require(len(values) == 1 and values[0].startswith("data:image/png;base64,"), "HTML page is not self-contained PNG")
        raw = base64.b64decode(values[0].split(",", 1)[1], validate=True)
        image = Image.open(io.BytesIO(raw))
        require(image.size == ((20, 10) if index == 1 else (10, 20)), "HTML raster dimensions changed")
        require(image.getpixel((0, 0)) == ((0, 0, 0, 255) if index == 1 else (40, 0, 0, 255)), "HTML raster order/samples changed")
        images.append({"sha256": hashlib.sha256(raw).hexdigest(), "size": list(image.size)})
    inferred = html.fromstring((root / "inferred-text.html").read_bytes())
    require(not inferred.xpath("//script|//img|//a|//form"), "Source text became active markup")
    require(inferred.xpath("//pre/text()") == ["</pre><script>x</script>&"], "Source literal text changed")
    return {"validator": "lxml HTML parser and Pillow PNG reader", "version": etree.LXML_VERSION, "offlineMarkupAndPixels": "passed", "embeddedImages": images,
            "actualBrowserAndAccessibilityQualification": "pending", "productionQualified": False,
            "artifacts": [{"file": name, "sha256": digest(root / name)} for name in ["fixed-pages.html", "inferred-text.html"]]}


def ebook_data(root):
    import zipfile
    import xml.etree.ElementTree as ET
    ns = {"o": "http://www.idpf.org/2007/opf", "h": "http://www.w3.org/1999/xhtml", "d": "http://purl.org/dc/elements/1.1/"}
    for name in ["fixed-book.epub", "reflow-book.epub"]:
        with zipfile.ZipFile(root / name) as package:
            first = package.infolist()[0]
            require(first.filename == "mimetype" and first.header_offset == 0 and first.compress_type == zipfile.ZIP_STORED, "EPUB stored first mimetype changed")
            require(package.read("mimetype") == b"application/epub+zip" and package.testzip() is None, "EPUB identity/CRC failed")
            opf = ET.fromstring(package.read("EPUB/package.opf"))
            require(opf.find("o:metadata/d:identifier", ns).text == "urn:uuid:11111111-2222-4333-8444-555555555555", "Explicit ebook identifier changed")
            require(opf.find("o:metadata/d:language", ns).text == "en", "Explicit ebook language changed")
            manifest = {item.get("id"): item.get("href") for item in opf.findall("o:manifest/o:item", ns)}
            spine = [manifest[item.get("idref")] for item in opf.findall("o:spine/o:itemref", ns)]
            nav = ET.fromstring(package.read("EPUB/nav.xhtml"))
            require([link.get("href") for link in nav.findall(".//h:nav/h:ol/h:li/h:a", ns)] == spine, "EPUB navigation/spine order changed")
            for index, path in enumerate(spine):
                chapter = ET.fromstring(package.read("EPUB/" + path))
                require(not chapter.findall(".//h:script", ns), "Source text became active ebook markup")
                if name == "fixed-book.epub":
                    image_path = chapter.find("h:body/h:img", ns).get("src")
                    image = Image.open(io.BytesIO(package.read("EPUB/" + image_path)))
                    require(image.size == ((20, 10) if index == 1 else (10, 20)), "EPUB raster page dimensions changed")
                    require(image.getpixel((0, 0)) == ((0, 0, 0, 255) if index == 1 else (40, 0, 0, 255)), "EPUB raster page order/samples changed")
                else:
                    require(chapter.find("h:body/h:pre", ns).text == "</pre><script>x</script>&", "EPUB inferred literal text changed")
    mobi = ET.parse(root / "book-source.xml")
    require(not mobi.findall(".//script"), "MobiXML source text became active markup")
    require(len(mobi.findall(".//{http://mobipocket.com/ns/mbp}pagebreak")) == 1, "MobiXML page break syntax changed")
    require([p.text for p in mobi.findall(".//pre")] == ["</pre><script>x</script>&"] * 2, "MobiXML literal text changed")
    return {"validators": ["Python zipfile/ElementTree", "Pillow"], "orderedPackageAndLiteralSourceData": "passed", "mobiCompilerAndCompleteReaderQualification": "pending", "productionQualified": False,
            "artifacts": [{"file": name, "sha256": digest(root / name)} for name in ["fixed-book.epub", "reflow-book.epub", "book-source.xml"]]}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("fixture_directory", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    if args.report.exists():
        previous = args.report.read_bytes()
        archived = args.report.with_name(args.report.stem + "-previous-" + hashlib.sha256(previous).hexdigest() + args.report.suffix)
        if archived.exists():
            require(archived.read_bytes() == previous, "Prior evidence archive differs")
        else:
            archived.write_bytes(previous)
        args.report.unlink()
    root = args.fixture_directory / "format-fixtures"
    paths = {name: root / name for name in ("vectors-source.svg", "vectors-export.svg", "vectors.pdf", "vectors-roundtrip.pdf")}
    source = Image.open(io.BytesIO(cairosvg.svg2png(bytestring=paths["vectors-source.svg"].read_bytes(), background_color="white"))).convert("RGB")
    exported = Image.open(io.BytesIO(cairosvg.svg2png(bytestring=paths["vectors-export.svg"].read_bytes(), background_color="white", output_width=200, output_height=120))).convert("RGB")
    original_pdf, roundtrip_pdf = pdf_pixels(paths["vectors.pdf"]), pdf_pixels(paths["vectors-roundtrip.pdf"])
    require(source.size == (200, 120) and exported.size == source.size and original_pdf.size == source.size and roundtrip_pdf.size == source.size, "Original physical SVG/PDF dimensions differ")
    svg_equal = ImageChops.difference(source, exported).getbbox() is None
    pdf_equal = ImageChops.difference(original_pdf, roundtrip_pdf).getbbox() is None
    require(svg_equal, "SVG export changed independently rendered original vectors/clips/styles")
    require(pdf_equal, "SVG round trip changed independently rendered PDF pixels")
    cross_engine_error = ImageStat.Stat(ImageChops.difference(source, original_pdf)).mean
    report = {"schemaVersion": 1, "purpose": "original-format-source-fixtures-not-production-qualification",
              "validators": {"CairoSVG": cairosvg.__version__, "PyMuPDF": fitz.VersionBind, "Pillow": PIL.__version__},
              "sourceSvgAndExportPixelsIdentical": svg_equal, "pdfRoundtripPixelsIdentical": pdf_equal,
              "crossRendererMeanAbsoluteRgbDifference": cross_engine_error,
              "crossRendererMetricScope": "Reported raw antialiasing difference; this is not a production visual-threshold qualification.",
              "artifacts": [{"file": name, "sha256": digest(path)} for name, path in paths.items()], "productionQualified": False}
    if (root / "table.xlsx").exists():
        report["spreadsheets"] = spreadsheet_data(root)
    if (root / "raster-slides.pptx").exists():
        report["presentations"] = presentation_data(root)
    if (root / "fixed-pages.html").exists():
        report["htmlExport"] = html_data(root)
    if (root / "fixed-book.epub").exists():
        report["ebooks"] = ebook_data(root)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2) + "\n")
    print("Independent controlled format fixture checks passed.")


if __name__ == "__main__":
    main()
